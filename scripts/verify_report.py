# -*- coding: utf-8 -*-
"""
verify_report.py — 输出报告验证
检查 zip 完整性 / XML合法性 / 全部写入单元格精确匹配 / 数据表行列 / 旧主题词残留 / 图表媒体保留。

用法:
  python verify_report.py --new 输出.xlsx --content content_data.py [--terms "复活节,Easter"] [--raw Product导出.xlsx]
"""
import argparse
import importlib.util
import re
import sys
import warnings
import zipfile
from xml.dom.minidom import parseString

import openpyxl


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--new', required=True)
    ap.add_argument('--content', required=True)
    ap.add_argument('--terms', default='', help='旧主题残留词, 逗号分隔')
    ap.add_argument('--raw', default=None)
    args = ap.parse_args()

    spec = importlib.util.spec_from_file_location('content_data', args.content)
    cd = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(cd)
    sheet_text = getattr(cd, 'SHEET_TEXT', {})
    data_sheet = getattr(cd, 'DATA_SHEET_RENAME', None) or getattr(cd, 'DATA_SHEET', None)

    fails = []

    # 1) zip + XML
    z = zipfile.ZipFile(args.new)
    bad = z.testzip()
    if bad:
        fails.append(f'zip损坏: {bad}')
    media_n = len([n for n in z.namelist() if n.startswith('xl/media/')])
    charts_n = len([n for n in z.namelist() if n.startswith('xl/charts/') and n.endswith('.xml')])
    for n in z.namelist():
        if n.endswith('.xml') or n.endswith('.rels'):
            try:
                parseString(z.read(n))
            except Exception as e:
                fails.append(f'XML语法错误 {n}: {e}')
    z.close()
    print(f'zip: {"OK" if not bad else "BAD"} | 媒体文件: {media_n} | 图表: {charts_n}')

    # 2) openpyxl 打开(openpyxl对WPS图表外部链接的警告为良性, 忽略)
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        wb = openpyxl.load_workbook(args.new)
    print('工作表:', wb.sheetnames)

    # 3) 逐单元格验证 SHEET_TEXT
    n_check = n_fail = 0
    for sheet_name, content in sheet_text.items():
        if sheet_name not in wb.sheetnames:
            fails.append(f'缺少工作表: {sheet_name}')
            continue
        ws = wb[sheet_name]
        for ref, want in content.items():
            got = ws[ref].value
            n_check += 1
            if isinstance(want, (int, float)) and not isinstance(want, bool):
                ok = isinstance(got, (int, float)) and abs(float(got) - float(want)) < 1e-9
            else:
                ok = str(got) == str(want)
            if not ok:
                n_fail += 1
                fails.append(f'[{sheet_name}!{ref}] 期望={str(want)[:60]!r} 实际={str(got)[:60]!r}')
    print(f'单元格验证: {n_check - n_fail}/{n_check} 通过')

    # 4) 数据表
    if data_sheet and data_sheet in wb.sheetnames:
        ws = wb[data_sheet]
        print(f'数据表[{data_sheet}]: {ws.max_row}行 x {ws.max_column}列')
        if args.raw:
            wbr = openpyxl.load_workbook(args.raw, read_only=True)
            wsr = wbr[wbr.sheetnames[0]]
            rows = list(wsr.iter_rows(values_only=True))
            wbr.close()
            expect_rows, expect_cols = len(rows), len(rows[0])
            if ws.max_row != expect_rows or ws.max_column != expect_cols:
                fails.append(f'数据表尺寸不符: {ws.max_row}x{ws.max_column} 期望 {expect_rows}x{expect_cols}')
            else:
                print(f'数据表尺寸匹配导出: {expect_rows}x{expect_cols}')
        if ws.max_row >= 2:
            print(f'  首行ASIN: {ws["B2"].value} | 标题: {str(ws["G2"].value or "")[:50]}')

    # 5) 旧主题残留(数据表除外)
    terms = [t.strip() for t in args.terms.split(',') if t.strip()]
    if terms:
        print(f'\n旧主题词残留扫描: {terms}')
        hits = []
        for sn in wb.sheetnames:
            if sn == data_sheet:
                continue
            ws = wb[sn]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and any(t in c.value for t in terms):
                        hits.append((sn, c.coordinate, c.value[:50]))
        if hits:
            print(f'发现{len(hits)}处(需人工判断是否有意保留的对比表述):')
            for h in hits[:20]:
                print('  ', h)
        else:
            print('无残留')

    wb.close()

    print()
    if fails:
        print(f'验证失败 {len(fails)}项:')
        for f_ in fails[:30]:
            print('  -', f_)
        sys.exit(1)
    print('验证全部通过 ✔')


if __name__ == '__main__':
    main()
