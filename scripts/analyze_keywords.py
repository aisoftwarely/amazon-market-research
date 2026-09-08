# -*- coding: utf-8 -*-
"""
analyze_keywords.py — 关键词挖掘 + 历史趋势 联合分析
挖掘表(每行一个关键词) join 历史表(每个sheet一个关键词的逐月序列),
计算季节性(峰/谷月/峰谷比)、近12月趋势、竞争度, 输出JSON+控制台摘要。

用法:
  python analyze_keywords.py <KeywordMining-*.xlsx> <KeywordHistory-*.xlsx> [--out keywords.json] [--top 20]
"""
import argparse
import json
import re
import openpyxl


def ppc_num(v):
    """'$0.95' -> 0.95"""
    if v is None or v == '':
        return None
    s = str(v).replace('$', '').strip()
    try:
        return float(s)
    except ValueError:
        return None


def load_mining(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h is not None else '' for h in rows[0]]
    idx = {h: i for i, h in enumerate(header)}
    need = ['关键词', '月搜索量', 'SPR', '需供比', '商品数', 'PPC竞价', '关键词翻译']
    missing = [k for k in need if k not in idx]
    if missing:
        print('警告: 挖掘表缺少关键列', missing, '| 实际表头:', header)
    kws = []
    for r in rows[1:]:
        if not r or not r[idx.get('关键词', 0)]:
            continue
        kws.append({
            'kw': str(r[idx['关键词']]).strip(),
            'translation': r[idx['关键词翻译']] if '关键词翻译' in idx else None,
            'volume': to_num(r[idx['月搜索量']]) if '月搜索量' in idx else None,
            'spr': to_num(r[idx['SPR']]) if 'SPR' in idx else None,
            'demand_supply': to_num(r[idx['需供比']]) if '需供比' in idx else None,
            'product_count': to_num(r[idx['商品数']]) if '商品数' in idx else None,
            'cpc': ppc_num(r[idx['PPC竞价']]) if 'PPC竞价' in idx else None,
            'top1_asin': r[idx['#1 前三ASIN']] if '#1 前三ASIN' in idx else None,
            'top1_click': to_num(r[idx['#1 点击共享']]) if '#1 点击共享' in idx else None,
            'top2_asin': r[idx['#2 前三ASIN']] if '#2 前三ASIN' in idx else None,
            'top2_click': to_num(r[idx['#2 点击共享']]) if '#2 点击共享' in idx else None,
            'top3_asin': r[idx['#3 前三ASIN']] if '#3 前三ASIN' in idx else None,
            'top3_click': to_num(r[idx['#3 点击共享']]) if '#3 点击共享' in idx else None,
        })
    return kws


def to_num(v):
    if v is None or v == '':
        return None
    try:
        f = float(v)
        return int(f) if f == int(f) else f
    except (TypeError, ValueError):
        return None


def load_history(path):
    """历史文件: 每个sheet一个关键词, 列含 月份/关键词/月搜索量"""
    wb = openpyxl.load_workbook(path, read_only=True)
    hist = {}
    for sn in wb.sheetnames:
        if 'note' in sn.lower():
            continue
        ws = wb[sn]
        rows = ws.iter_rows(values_only=True)
        try:
            header = [str(h).strip() if h is not None else '' for h in next(rows)]
        except StopIteration:
            continue
        if '月份' not in header or '月搜索量' not in header:
            continue
        i_m, i_v = header.index('月份'), header.index('月搜索量')
        i_kw = header.index('关键词') if '关键词' in header else None
        series = []
        kw_name = None
        for r in rows:
            if not r or r[i_m] is None:
                continue
            if i_kw is not None and r[i_kw]:
                kw_name = str(r[i_kw]).strip()
            series.append((str(r[i_m])[:7], to_num(r[i_v])))
        if not series:
            continue
        # sheet名兜底提取关键词: History-bride socks-US -> bride socks
        m = re.match(r'History-(.*?)-US\d*$', sn)
        sheet_kw = m.group(1) if m else None
        key = kw_name or sheet_kw or sn
        series.sort(key=lambda t: t[0])
        hist[key] = {'series': series, 'sheet': sn, 'sheet_kw': sheet_kw}
    wb.close()
    return hist


def analyze_series(series):
    """[(YYYY-MM, vol), ...] -> 季节性指标"""
    if not series or len(series) < 6:
        return None
    vols = [(m, v) for m, v in series if v is not None]
    if len(vols) < 6:
        return None
    peak = max(vols, key=lambda t: t[1])
    trough = min(vols, key=lambda t: t[1])
    # 各月份历史均值(跨年), 判断旺季月份
    month_sum = {}
    for m, v in vols:
        mm = m[5:7]
        month_sum.setdefault(mm, []).append(v)
    month_avg = {mm: sum(vs) / len(vs) for mm, vs in month_sum.items()}
    overall = sum(v for _, v in vols) / len(vols)
    peak_months = sorted([mm for mm, a in month_avg.items() if a >= overall * 1.15])
    # 近12月 vs 前12月
    recent = vols[-12:]
    prior = vols[-24:-12]
    recent_avg = sum(v for _, v in recent) / len(recent)
    prior_avg = (sum(v for _, v in prior) / len(prior)) if prior else None
    yoy = (recent_avg / prior_avg - 1) if prior and prior_avg else None
    return {
        'first_month': vols[0][0], 'last_month': vols[-1][0], 'last_volume': vols[-1][1],
        'peak_month': peak[0], 'peak_volume': peak[1], 'trough_month': trough[0], 'trough_volume': trough[1],
        'peak_trough_ratio': round(peak[1] / trough[1], 2) if trough[1] else None,
        'high_months': peak_months, 'recent_12m_avg': round(recent_avg, 1),
        'yoy_trend': round(yoy, 3) if yoy is not None else None,
        'series': [{'month': m, 'vol': v} for m, v in vols],
    }


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('mining_file')
    ap.add_argument('history_file')
    ap.add_argument('--out', default=None)
    ap.add_argument('--top', type=int, default=20)
    args = ap.parse_args()

    kws = load_mining(args.mining_file)
    hist = load_history(args.history_file)
    print(f'挖掘表: {len(kws)}个关键词 | 历史表: {len(hist)}个关键词sheet')

    # join: 精确匹配 -> 前缀匹配(sheet名被截断)
    hist_by_prefix = {}
    for k, v in hist.items():
        hist_by_prefix[k] = v
    joined = 0
    results = []
    for k in kws:
        h = hist.get(k['kw'])
        if not h:
            for hk, hv in hist.items():
                sk = hv.get('sheet_kw') or hk
                if sk and (k['kw'].startswith(sk) or sk.startswith(k['kw'])) and len(sk) >= 6:
                    h = hv
                    break
        item = dict(k)
        if h:
            joined += 1
            sa = analyze_series(h['series'])
            if sa:
                item['trend'] = sa
        results.append(item)
    print(f'成功join历史趋势: {joined}/{len(kws)}')

    results.sort(key=lambda x: -(x.get('volume') or 0))
    print(f'\n===== Top{args.top} 关键词 (按月搜索量) =====')
    print(f'{"关键词":<42}{"月搜":>9}{"CPC":>6}{"SPR":>5}{"需供比":>7}  季节性')
    for r in results[:args.top]:
        t = r.get('trend')
        if t and t.get('peak_trough_ratio'):
            season = f"峰{t['peak_month']}({int(t['peak_volume'])})/谷{t['trough_month']}({int(t['trough_volume'])}) 峰谷比{t['peak_trough_ratio']}"
        else:
            season = '-'
        spr = r.get('spr') if r.get('spr') is not None else '-'
        ds = r.get('demand_supply') if r.get('demand_supply') is not None else '-'
        print(f"{r['kw']:<42}{int(r.get('volume') or 0):>9,}{r.get('cpc') or 0:>6.2f}{spr!s:>5}{ds!s:>7}  {season}")

    if args.out:
        with open(args.out, 'w', encoding='utf-8') as f:
            json.dump(results, f, ensure_ascii=False, indent=1, default=str)
        print(f'\nJSON已输出 -> {args.out}')


if __name__ == '__main__':
    main()
