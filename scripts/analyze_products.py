# -*- coding: utf-8 -*-
"""
analyze_products.py — 卖家精灵 Product 商品导出分析
计算市场统计/价格带/品牌集中度/评论门槛/新品友好度/去重母体Top竞品, 输出JSON+控制台摘要。

用法:
  python analyze_products.py <Product导出.xlsx> [--out products.json] [--bands 0,8,12,15,25] [--top 7]

列名自动按表头发现, 依赖卖家精灵标准中文表头(ASIN/品牌/月销量/价格($)/评分数/上架天数/父ASIN等)。
"""
import argparse
import json
import datetime
import openpyxl

# 表头 -> 内部键 (卖家精灵Product导出标准表头)
COLMAP = {
    'ASIN': 'asin', '父ASIN': 'parent', '品牌': 'brand', '商品标题': 'title',
    '类目路径': 'cat', '大类目': 'cat_main', '大类BSR': 'bsr', '小类BSR': 'sub_bsr',
    '月销量': 'sales', '销量环比增长率': 'mom', '销量同比增长率': 'yoy',
    '月销售额($)': 'revenue', '变体数': 'variants', '价格($)': 'price',
    '评分数': 'rating_count', '月新增评分数': 'new_reviews', '评分': 'rating',
    'FBA($)': 'fba', '毛利率': 'margin', '上架时间': 'launch', '上架天数': 'days',
    '配送方式': 'fulfillment', 'A+页面': 'a_plus', '视频介绍': 'video',
    '包装尺寸分段': 'size_tier', '卖家所属地': 'seller_region', 'BuyBox卖家': 'buybox',
    '商品重量': 'weight', '包装重量': 'pkg_weight', '包装尺寸': 'pkg_size',
    '详细参数': 'params', '小类目': 'cat_sub',
}

NUM_KEYS = {'sales', 'revenue', 'variants', 'price', 'rating_count', 'new_reviews',
            'rating', 'fba', 'margin', 'days', 'bsr', 'mom', 'yoy', 'sub_bsr'}


def to_num(v):
    if v is None or v == '':
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(',', '').replace('%', '')
    try:
        f = float(s)
        return int(f) if f == int(f) else f
    except ValueError:
        return v


def load_products(path):
    wb = openpyxl.load_workbook(path, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h is not None else '' for h in rows[0]]
    idx = {}
    for i, h in enumerate(header):
        if h in COLMAP:
            idx[COLMAP[h]] = i
    prods = []
    for r in rows[1:]:
        if r is None or all(v in (None, '') for v in r):
            continue
        p = {}
        for k, i in idx.items():
            v = r[i] if i < len(r) else None
            p[k] = to_num(v) if k in NUM_KEYS else v
        if isinstance(p.get('launch'), (datetime.date, datetime.datetime)):
            p['launch'] = str(p['launch'])[:10]
        if p.get('asin'):
            prods.append(p)
    return prods, idx


def fnum(v, nd=0):
    if v is None:
        return '-'
    return f'{v:,.{nd}f}'


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('product_export')
    ap.add_argument('--out', default=None)
    ap.add_argument('--bands', default='0,8,12,15,25')
    ap.add_argument('--top', type=int, default=7)
    args = ap.parse_args()

    prods, idx = load_products(args.product_export)
    n = len(prods)
    missing = [k for k in ['asin', 'sales', 'price'] if k not in idx]
    if missing:
        print(f'警告: 缺少关键列 {missing}, 请核对导出文件表头')
        print('发现表头映射:', {k: i for k, i in idx.items()})

    total_sales = sum(p['sales'] or 0 for p in prods if p.get('sales'))
    total_rev = sum(p['revenue'] or 0 for p in prods if p.get('revenue'))
    asp = total_rev / total_sales if total_sales else None
    avg_rating = [p['rating'] for p in prods if p.get('rating')]
    avg_rating = sum(avg_rating) / len(avg_rating) if avg_rating else None

    print(f'===== 市场总览 (样本 {n} ASIN) =====')
    print(f'月总销量: {fnum(total_sales)}件 | 月总销售额: ${fnum(total_rev)} | ASP: ${asp:.2f}' if asp else '无销售额数据')
    if avg_rating:
        print(f'均评分: {avg_rating:.2f} | 单ASIN平均月销: {fnum(total_sales / n)}件')

    # ---- 价格带 ----
    edges = [float(x) for x in args.bands.split(',')]
    bands = []
    print(f'\n===== 价格带 ({"-".join(str(int(e)) for e in edges)}) =====')
    for lo, hi in zip(edges, edges[1:]):
        grp = [p for p in prods if p.get('price') is not None and lo <= p['price'] < hi]
        bs = sum(p['sales'] or 0 for p in grp)
        br = sum(p['revenue'] or 0 for p in grp)
        rep = sorted(grp, key=lambda p: -(p['sales'] or 0))[:3]
        rep_s = '; '.join(f"${p['price']} {p['brand']}(月销{int(p['sales'])})" for p in rep if p.get('sales'))
        line = (f'${lo:g}-{hi:g}: {len(grp)}款({len(grp)/n:.1%}) 销量{int(bs)}件({bs/total_sales:.1%}) '
                f'销售额${int(br):,}({br/total_rev:.1%}) | 代表: {rep_s}')
        print(line)
        bands.append({'band': f'${lo:g}-{hi:g}', 'count': len(grp), 'share': round(len(grp) / n, 4),
                      'sales': int(bs), 'sales_share': round(bs / total_sales, 4) if total_sales else None,
                      'revenue': round(br, 2), 'revenue_share': round(br / total_rev, 4) if total_rev else None,
                      'representatives': [{'asin': p['asin'], 'brand': p.get('brand'), 'price': p.get('price'),
                                           'sales': p.get('sales')} for p in rep]})

    # ---- 品牌集中度 ----
    brand_sales = {}
    for p in prods:
        if p.get('sales'):
            brand_sales[p.get('brand') or '(无品牌)'] = brand_sales.get(p.get('brand') or '(无品牌)', 0) + p['sales']
    top_brands = sorted(brand_sales.items(), key=lambda kv: -kv[1])[:5]
    print(f'\n===== Top5品牌 (销量集中度) =====')
    for b, s in top_brands:
        print(f'{b}: {int(s)}件 ({s/total_sales:.1%})')

    # ---- 评论门槛 ----
    rc_bands = [(0, 50), (50, 100), (100, 500), (500, 10**9)]
    print(f'\n===== 评论门槛 (销量按评分数分布) =====')
    review_bands = []
    for lo, hi in rc_bands:
        grp = [p for p in prods if p.get('rating_count') is not None and lo <= p['rating_count'] < hi]
        s = sum(p['sales'] or 0 for p in grp)
        print(f'{lo}+评: {len(grp)}款, 销量占比 {s/total_sales:.1%}' if total_sales else f'{lo}+评: {len(grp)}款')
        review_bands.append({'band': f'{lo}+', 'count': len(grp), 'sales': int(s),
                             'sales_share': round(s / total_sales, 4) if total_sales else None})

    # ---- 新品友好度 ----
    NEW_DAYS = 456  # 约15个月
    new_grp = [p for p in prods if p.get('days') is not None and p['days'] <= NEW_DAYS]
    ns = sum(p['sales'] or 0 for p in new_grp)
    print(f'\n===== 新品友好度 =====')
    print(f'上架≤15个月: {len(new_grp)}款({len(new_grp)/n:.1%}), 销量占比 {ns/total_sales:.1%}')
    top_sales = sorted(prods, key=lambda p: -(p['sales'] or 0))
    if top_sales:
        t = top_sales[0]
        print(f"销量Top1: {t['asin']} {t.get('brand')} ${t.get('price')} 月销{int(t.get('sales') or 0)} "
              f"{t.get('rating')}星{int(t.get('rating_count') or 0)}评")

    # ---- Top10集中度 ----
    top10 = sum(p['sales'] or 0 for p in top_sales[:10])
    print(f'Top10销量集中度: {top10/total_sales:.1%}')

    # ---- 去重母体(父ASIN聚合) ----
    parents = {}
    for p in prods:
        key = p.get('parent') or p['asin']
        d = parents.setdefault(key, {'parent': key, 'children': [], 'sales': 0, 'revenue': 0})
        d['children'].append(p)
        d['sales'] += p['sales'] or 0
        d['revenue'] += p['revenue'] or 0
    parent_list = sorted(parents.values(), key=lambda d: -d['sales'])
    print(f'\n===== 去重母体: {len(parent_list)}个 (Top{args.top}) =====')
    top_parents = []
    for d in parent_list[:args.top]:
        best = sorted(d['children'], key=lambda p: -(p['sales'] or 0))[0]
        top_parents.append({
            'parent': d['parent'], 'parent_sales': int(d['sales']),
            'best_child': best, 'children_count': len(d['children']),
        })
        b = best
        print(f"--- {b['asin']} ({b.get('brand')}) ${b.get('price')} 母体月销{int(d['sales'])} "
              f"(子体{len(d['children'])}) {b.get('rating')}星{int(b.get('rating_count') or 0)}评 "
              f"上架{b.get('launch')} 环比{b.get('mom')} FBA{b.get('fba')} 毛利率{b.get('margin')}")
        print(f"    TITLE: {b.get('title')}")

    # ---- 利润参考 ----
    fba_list = [p['fba'] for p in prods if p.get('fba') is not None]
    mg_list = [p['margin'] for p in prods if isinstance(p.get('margin'), (int, float))]
    print(f'\n===== 利润参考 =====')
    if fba_list:
        print(f'FBA费范围: ${min(fba_list):.2f} - ${max(fba_list):.2f}')
    if mg_list:
        print(f'毛利率范围: {min(mg_list):.1%} - {max(mg_list):.1%}')

    # ---- 输出JSON ----
    if args.out:
        result = {
            'sample_size': n, 'total_sales': int(total_sales), 'total_revenue': round(total_rev, 2),
            'asp': round(asp, 2) if asp else None, 'avg_rating': round(avg_rating, 2) if avg_rating else None,
            'price_bands': bands, 'top_brands': [{'brand': b, 'sales': int(s), 'share': round(s / total_sales, 4)}
                                                 for b, s in top_brands],
            'review_bands': review_bands,
            'new_product_sales_share': round(ns / total_sales, 4) if total_sales else None,
            'top10_sales_share': round(top10 / total_sales, 4) if total_sales else None,
            'top_parents': top_parents,
            'all_products': prods,
        }
        with open(args.out, 'w', encoding='utf-8') as f:
            json.dump(result, f, ensure_ascii=False, indent=1, default=str)
        print(f'\nJSON已输出 -> {args.out}')


if __name__ == '__main__':
    main()
