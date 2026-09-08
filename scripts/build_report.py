# -*- coding: utf-8 -*-
"""
build_report.py — Excel 调研模板 XML 直改写入引擎
在不破坏模板样式/嵌入图片(WPS DISPIMG)/图表/超链接的前提下, 将分析内容写入模板。

核心原理: xlsx = zip 包。解包全部 XML -> 按规则修改 -> 全量语法校验 -> 重新打包。
绝不使用 openpyxl 保存(openpyxl 会丢弃 WPS DISPIMG 单元格图片)。

用法:
  python build_report.py --template 模板.xlsx --content content_data.py \
      --raw Product-US-导出.xlsx --out 输出.xlsx

content_data.py 需定义的配置见 assets/content_data_template.py, 摘要:
  SHEET_TEXT        {工作表名: {单元格ref: 值}}            文本/数字写入(保留样式)
  DONOR_STYLES      {工作表名: {ref: 同表donor ref}}        新单元格样式继承
  DATA_SHEET        '原始数据表名'                          整表重建为本次导出数据
  DATA_SHEET_RENAME '新表名'                               数据表改名(年月标识)
  RAW_LINK_COLS     ['ASIN','品牌链接',...]                 数据表自动超链接列(按表头)
  CHART_UPDATES     {图表xml: {text_replaces, series}}     图表缓存更新
  SHEET_LINKS       {工作表名: {ref: url}}                  既有表超链接目标更新
  RAW_PATCHES       {xml路径: [(old, new)]}                任意XML字面替换(逃生舱)
  MEDIA_ADD         {zip内媒体路径: 本地图片}               新增媒体文件
  DELETE_ANCHOR_RIDS {drawing.xml: [rIdN]}                 删除图片锚点
"""
import argparse
import datetime
import importlib.util
import os
import re
import sys
import zipfile
from xml.dom.minidom import parseString

import openpyxl

HL_REL_TYPE = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink'


def esc(t):
    return t.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;').replace('"', '&quot;')


# ============ 列号工具 ============
def col_to_num(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def num_to_col(n):
    s = ''
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


class Builder:
    def __init__(self, template_path, content_mod, raw_path, out_path):
        self.out = out_path
        self.raw_path = raw_path
        z = zipfile.ZipFile(template_path)
        self.names = z.namelist()
        self.parts = {n: z.read(n) for n in self.names}
        z.close()

        g = lambda k, d=None: getattr(content_mod, k, d)
        self.sheet_text = g('SHEET_TEXT', {})
        self.donor_styles = g('DONOR_STYLES', {})
        self.data_sheet = g('DATA_SHEET')
        self.data_sheet_rename = g('DATA_SHEET_RENAME') or self.data_sheet
        self.raw_link_cols = g('RAW_LINK_COLS', ['ASIN', '品牌链接', '商品详情页链接', '商品主图', '父ASIN', '卖家首页'])
        self.chart_updates = g('CHART_UPDATES', {})
        self.sheet_links = g('SHEET_LINKS', {})
        self.raw_patches = g('RAW_PATCHES', {})
        self.media_add = g('MEDIA_ADD', {})
        self.delete_anchor_rids = g('DELETE_ANCHOR_RIDS', {})

        # 工作表名 -> 文件路径
        wbx = self.parts['xl/workbook.xml'].decode('utf-8')
        wbrels = self.parts['xl/_rels/workbook.xml.rels'].decode('utf-8')
        rid2target = dict(re.findall(r'<Relationship Id="(rId\d+)"[^>]*Target="([^"]+)"', wbrels))
        self.sheet_files = {}
        for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx):
            t = rid2target.get(m.group(2), '')
            if not t.startswith('xl/'):
                t = 'xl/' + t.lstrip('/')
            self.sheet_files[m.group(1)] = t

        # sharedStrings: 原条目原样保留(自带<t>包裹!), 新条目才需要包<t>
        ss = self.parts['xl/sharedStrings.xml'].decode('utf-8')
        self.strings = re.findall(r'<si>(.*?)</si>', ss, re.S)
        self.n_orig = len(self.strings)
        self.ss_map = {}
        for i, s in enumerate(self.strings):
            self.ss_map.setdefault(s, i)

    # ============ sharedStrings ============
    def si_idx(self, text):
        raw = esc(text)
        if raw in self.ss_map:
            return self.ss_map[raw]
        idx = len(self.strings)
        self.strings.append(raw)
        self.ss_map[raw] = idx
        return idx

    def rebuild_sharedstrings(self):
        total = 0
        for n in self.names:
            if n.startswith('xl/worksheets/sheet') and n.endswith('.xml'):
                total += len(re.findall(r't="s"', self.parts[n].decode('utf-8')))
        ss = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
              f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
              f'count="{total}" uniqueCount="{len(self.strings)}">')
        for i, raw in enumerate(self.strings):
            if i < self.n_orig:
                ss += f'<si>{raw}</si>'  # 原样保留, 防止双包裹导致全部文本丢失
            else:
                attr = ' xml:space="preserve"' if raw != raw.strip() else ''
                ss += f'<si><t{attr}>{raw}</t></si>'
        ss += '</sst>'
        self.parts['xl/sharedStrings.xml'] = ss.encode('utf-8')
        print(f'sharedStrings: {len(self.strings)}条(原{self.n_orig}+新{len(self.strings) - self.n_orig}), {total}引用')

    # ============ 单元格引擎 ============
    def make_cell(self, ref, value, style=None):
        sattr = f' s="{style}"' if style is not None else ''
        if value is None:
            return f'<c r="{ref}"{sattr}/>'
        if isinstance(value, bool):
            return f'<c r="{ref}"{sattr} t="b"><v>{1 if value else 0}</v></c>'
        if isinstance(value, datetime.datetime):
            value = value.strftime('%Y-%m-%d %H:%M:%S')
        if isinstance(value, datetime.date):
            value = value.strftime('%Y-%m-%d')
        if isinstance(value, (int, float)):
            v = value
            if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
                v = int(v)
            return f'<c r="{ref}"{sattr}><v>{v}</v></c>'
        idx = self.si_idx(value)
        return f'<c r="{ref}"{sattr} t="s"><v>{idx}</v></c>'

    def set_cell(self, xml, ref, value, donor_style=None):
        m = re.match(r'([A-Z]+)(\d+)', ref)
        col, row, rn = m.group(1), m.group(2), int(m.group(2))
        row_m = re.search(r'<row r="%s"([^>]*?)(/>|>)' % row, xml)
        if not row_m:
            # 行不存在: 按行号顺序插入新行
            new_row = f'<row r="{row}">{self.make_cell(ref, value, donor_style)}</row>'
            pos = None
            for rm in re.finditer(r'<row r="(\d+)"', xml):
                if int(rm.group(1)) > rn:
                    pos = rm.start()
                    break
            if pos is None:
                pos = xml.find('</sheetData>')
                if pos < 0:
                    pos = xml.find('</worksheet>')
            return xml[:pos] + new_row + xml[pos:]
        attrs = row_m.group(1)
        if row_m.group(2) == '/>':
            # 自闭合空行 -> 展开
            new_row = f'<row r="{row}"{attrs}>{self.make_cell(ref, value, donor_style)}</row>'
            return xml[:row_m.start()] + new_row + xml[row_m.end():]
        content_start = row_m.end()
        close = xml.find('</row>', content_start)
        row_content = xml[content_start:close]
        cell_m = re.search(r'<c r="%s"([^>]*?)(/>|>.*?</c>)' % re.escape(ref), row_content, re.S)
        if cell_m:
            sm = re.search(r's="(\d+)"', cell_m.group(1))
            style = sm.group(1) if sm else donor_style
            new_cell = self.make_cell(ref, value, style)
            row_content = row_content[:cell_m.start()] + new_cell + row_content[cell_m.end():]
        else:
            new_cell = self.make_cell(ref, value, donor_style)
            existing = sorted([(col_to_num(cm.group(1)), cm.start()) for cm in
                               re.finditer(r'<c r="([A-Z]+)\d+"', row_content)])
            target = col_to_num(col)
            pos = len(row_content)
            for cn, st in existing:
                if cn > target:
                    pos = st
                    break
            row_content = row_content[:pos] + new_cell + row_content[pos:]
        return xml[:content_start] + row_content + xml[close:]

    def apply_sheet_text(self):
        for sheet_name, content in self.sheet_text.items():
            f = self.sheet_files[sheet_name]
            xml = self.parts[f].decode('utf-8')
            donors = self.donor_styles.get(sheet_name, {})
            for ref, value in content.items():
                donor = None
                dref = donors.get(ref)
                if dref:
                    dm = re.search(r'<c r="%s"[^>]*?s="(\d+)"' % re.escape(dref), xml)
                    if dm:
                        donor = dm.group(1)
                xml = self.set_cell(xml, ref, value, donor)
            self.parts[f] = xml.encode('utf-8')
        print(f'文本写入: {sum(len(c) for c in self.sheet_text.values())}个单元格, {len(self.sheet_text)}个表')

    # ============ 数据表重建 ============
    @staticmethod
    def conv(v):
        if v is None:
            return None
        if isinstance(v, (int, float)):
            if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
                return int(v)
            return v
        if isinstance(v, str):
            s = v.strip()
            if not s:
                return None
            if re.match(r'^-?\d+(\.\d+)?$', s):
                f = float(s)
                return int(f) if f == int(f) else f
            return v
        return v

    def rebuild_data_sheet(self):
        if not self.data_sheet:
            return
        f = self.sheet_files[self.data_sheet]
        xml = self.parts[f].decode('utf-8')

        wb = openpyxl.load_workbook(self.raw_path, read_only=True)
        ws = wb[wb.sheetnames[0]]
        raw_rows = list(ws.iter_rows(values_only=True))
        wb.close()
        header = [str(h).strip() if h is not None else '' for h in raw_rows[0]]
        data = raw_rows[1:]
        ncols = len(header)

        # 列样式: 沿用模板原row1/row2/row3
        col_styles = {}
        for rn in ['1', '2', '3']:
            rm = re.search(r'<row r="%s"[^>]*>(.*?)</row>' % rn, xml, re.S)
            if not rm:
                continue
            for cm in re.finditer(r'<c r="([A-Z]+)%s"([^>]*?)(?:/>|>)' % rn, rm.group(1)):
                ci = col_to_num(cm.group(1)) - 1
                sm = re.search(r's="(\d+)"', cm.group(2))
                if sm and ci not in col_styles:
                    col_styles[ci] = sm.group(1)
        row2m = re.search(r'<row r="2"([^>]*?)>', xml)
        row_attrs = row2m.group(1) if row2m else ''

        rows_xml = []
        for ri_minus, r in enumerate([header] + data):
            ri = ri_minus + 1
            cells = []
            for ci in range(ncols):
                ref = f'{num_to_col(ci + 1)}{ri}'
                v = r[ci] if ci < len(r) else None
                if isinstance(v, (datetime.date, datetime.datetime)):
                    v = str(v)[:10]
                cells.append(self.make_cell(ref, self.conv(v), col_styles.get(ci)))
            rows_xml.append(f'<row r="{ri}"{row_attrs}>' + ''.join(cells) + '</row>')
        new_sd = '<sheetData>' + ''.join(rows_xml) + '</sheetData>'
        xml = re.sub(r'<sheetData>.*?</sheetData>', new_sd, xml, flags=re.S)
        xml = re.sub(r'<dimension ref="[^"]*"/>',
                     f'<dimension ref="A1:{num_to_col(ncols)}{len(data) + 1}"/>', xml)

        # ---- 自动超链接(按表头名) ----
        link_col_idx = {}
        for col_name in self.raw_link_cols:
            if col_name in header:
                link_col_idx[header.index(col_name)] = col_name
        hl_entries, rel_entries = [], []
        rid_n = 0
        for i, r in enumerate(data):
            ri = i + 2
            for ci, col_name in link_col_idx.items():
                v = r[ci] if ci < len(r) else None
                if not v:
                    continue
                v = str(v).strip()
                if col_name in ('ASIN', '父ASIN'):
                    target = f'https://www.amazon.com/dp/{v}'
                else:
                    target = v
                rid_n += 1
                rid = f'rId{rid_n}'
                ref = f'{num_to_col(ci + 1)}{ri}'
                hl_entries.append(f'<hyperlink ref="{ref}" r:id="{rid}" display="{esc(v)[:255]}"/>')
                rel_entries.append(f'<Relationship Id="{rid}" Type="{HL_REL_TYPE}" '
                                   f'Target="{esc(target)[:500]}" TargetMode="External"/>')
        if hl_entries:
            if '<hyperlinks>' in xml:
                xml = re.sub(r'<hyperlinks>.*?</hyperlinks>',
                             '<hyperlinks>' + ''.join(hl_entries) + '</hyperlinks>', xml, flags=re.S)
            else:
                dm = re.search(r'<drawing ', xml)
                pos = dm.start() if dm else xml.find('</worksheet>')
                xml = xml[:pos] + '<hyperlinks>' + ''.join(hl_entries) + '</hyperlinks>' + xml[pos:]

        self.parts[f] = xml.encode('utf-8')
        relpath = 'xl/worksheets/_rels/' + os.path.basename(f) + '.rels'
        self.parts[relpath] = (('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
                                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                                + ''.join(rel_entries) + '</Relationships>')).encode('utf-8')
        print(f'数据表重建[{self.data_sheet}]: {len(data)}行x{ncols}列, {len(rel_entries)}超链接')

        # ---- 改名 ----
        if self.data_sheet_rename and self.data_sheet_rename != self.data_sheet:
            for part in ['xl/workbook.xml', 'docProps/app.xml']:
                if part in self.parts:
                    x = self.parts[part].decode('utf-8')
                    x = x.replace(self.data_sheet, self.data_sheet_rename)
                    self.parts[part] = x.encode('utf-8')
            print(f'数据表改名: {self.data_sheet} -> {self.data_sheet_rename}')

    # ============ 图表更新 ============
    def update_charts(self):
        for path, cfg in self.chart_updates.items():
            if path not in self.parts:
                print(f'警告: 图表不存在 {path}')
                continue
            chart = self.parts[path].decode('utf-8')
            for old, new in cfg.get('text_replaces', []):
                chart = chart.replace(old, new)
            for series in cfg.get('series', []):
                cat_zone = re.search(r'<c:cat>.*?</c:cat>', chart, re.S)
                val_zone = re.search(r'<c:val>.*?</c:val>', chart, re.S)
                if cat_zone:
                    chart = chart.replace(cat_zone.group(0), self._rebuild_pts(cat_zone.group(0), series['cats']))
                if val_zone:
                    chart = chart.replace(val_zone.group(0), self._rebuild_pts(val_zone.group(0), series['vals']))
            parseString(chart.encode('utf-8'))  # 语法自检
            self.parts[path] = chart.encode('utf-8')
            print(f'图表更新: {path}')

    @staticmethod
    def _rebuild_pts(zone, items):
        # 只替换cache内的pt序列, 保留f引用/结构 (整段重建会破坏结构)
        zone = re.sub(r'<c:pt idx="\d+">.*?</c:pt>', '', zone, flags=re.S)
        pts = ''.join(f'<c:pt idx="{i}"><c:v>{esc(str(it))}</c:v></c:pt>' for i, it in enumerate(items))
        return re.sub(r'<c:ptCount val="\d+"/>', f'<c:ptCount val="{len(items)}"/>' + pts, zone)

    # ============ 工作表超链接更新 ============
    def update_sheet_links(self):
        for sheet_name, links in self.sheet_links.items():
            f = self.sheet_files[sheet_name]
            xml = self.parts[f].decode('utf-8')
            relpath = 'xl/worksheets/_rels/' + os.path.basename(f) + '.rels'
            if relpath in self.parts:
                rels = self.parts[relpath].decode('utf-8')
            else:
                rels = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\r\n'
                        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"></Relationships>')
            hl_m = re.search(r'<hyperlinks>(.*?)</hyperlinks>', xml, re.S)
            entries = re.findall(r'<hyperlink[^>]*/>', hl_m.group(1)) if hl_m else []
            ref2rid = {}
            kept = []
            for e in entries:
                rm = re.search(r'ref="([A-Z]+\d+)" r:id="(rId\d+)"', e)
                if not rm:
                    continue
                ref2rid[rm.group(1)] = rm.group(2)
                if rm.group(1) not in links:
                    kept.append(e)
            used = [int(r[3:]) for r in re.findall(r'Id="(rId\d+)"', rels)]
            next_n = (max(used) + 1) if used else 1
            new_entries = []
            for ref, url in links.items():
                if ref in ref2rid:
                    rid = ref2rid[ref]
                    rels = re.sub(r'(<Relationship Id="%s"[^>]*?Target=")[^"]*(")' % rid,
                                  lambda m: m.group(1) + esc(url)[:500] + m.group(2), rels)
                else:
                    rid = f'rId{next_n}'
                    next_n += 1
                    rels += (f'<Relationship Id="{rid}" Type="{HL_REL_TYPE}" '
                             f'Target="{esc(url)[:500]}" TargetMode="External"/>')
                new_entries.append(f'<hyperlink ref="{ref}" r:id="{rid}" display="{esc(url)[:255]}"/>')
            block = '<hyperlinks>' + ''.join(kept + new_entries) + '</hyperlinks>'
            if hl_m:
                xml = xml[:hl_m.start()] + block + xml[hl_m.end():]
            else:
                dm = re.search(r'<drawing ', xml)
                pos = dm.start() if dm else xml.find('</worksheet>')
                xml = xml[:pos] + block + xml[pos:]
            self.parts[f] = xml.encode('utf-8')
            self.parts[relpath] = rels.encode('utf-8')
            print(f'表超链接更新[{sheet_name}]: {len(links)}个')

    # ============ 任意XML字面替换 ============
    def apply_raw_patches(self):
        for path, pairs in self.raw_patches.items():
            if path not in self.parts:
                print(f'警告: RAW_PATCHES目标不存在 {path}')
                continue
            x = self.parts[path].decode('utf-8')
            for old, new in pairs:
                if old not in x:
                    print(f'警告: RAW_PATCHES未命中 [{path}] {old[:50]!r}')
                x = x.replace(old, new)
            self.parts[path] = x.encode('utf-8')
        if self.raw_patches:
            print(f'RAW_PATCHES: {len(self.raw_patches)}个文件')

    # ============ 媒体与图片锚点 ============
    def apply_media(self):
        for zip_path, local_path in self.media_add.items():
            with open(local_path, 'rb') as fh:
                self.parts[zip_path] = fh.read()
            print(f'新增媒体: {zip_path} <- {local_path}')
        for dpath, rids in self.delete_anchor_rids.items():
            if dpath not in self.parts:
                continue
            d = self.parts[dpath].decode('utf-8')
            for rid in rids:
                d = re.sub(r'<xdr:(?:twoCell|oneCell)Anchor[^>]*>(?:(?!</xdr:(?:twoCell|oneCell)Anchor>).)*?'
                           r'r:embed="%s"(?:(?!</xdr:(?:twoCell|oneCell)Anchor>).)*?'
                           r'</xdr:(?:twoCell|oneCell)Anchor>' % rid, '', d, flags=re.S)
            self.parts[dpath] = d.encode('utf-8')
            n_left = len(re.findall(r'<xdr:(?:twoCell|oneCell)Anchor', d))
            print(f'图片锚点清理[{dpath}]: 剩余{n_left}个锚点')

    # ============ 校验与写出 ============
    def validate_all(self):
        for n in list(self.parts):
            if n.endswith('.xml') or n.endswith('.rels'):
                try:
                    parseString(self.parts[n])
                except Exception as e:
                    print(f'XML语法错误 {n}: {e}')
                    sys.exit(1)
        print('全量XML语法校验通过')

    def write(self):
        with zipfile.ZipFile(self.out, 'w', zipfile.ZIP_DEFLATED) as zo:
            written = set()
            for n in self.names:  # 原顺序(含[Content_Types].xml居首)
                zo.writestr(n, self.parts[n])
                written.add(n)
            for n in self.parts:  # 新增部件: rels/媒体等
                if n not in written:
                    zo.writestr(n, self.parts[n])
                    print(f'新增zip部件: {n}')
        print(f'写出: {self.out} ({os.path.getsize(self.out):,} bytes)')

    def build(self):
        self.apply_sheet_text()
        self.rebuild_data_sheet()
        self.update_charts()
        self.update_sheet_links()
        self.apply_raw_patches()
        self.apply_media()
        self.rebuild_sharedstrings()
        self.validate_all()
        self.write()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--template', required=True)
    ap.add_argument('--content', required=True)
    ap.add_argument('--raw', required=True)
    ap.add_argument('--out', required=True)
    args = ap.parse_args()

    spec = importlib.util.spec_from_file_location('content_data', args.content)
    cd = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(cd)

    Builder(args.template, cd, args.raw, args.out).build()


if __name__ == '__main__':
    main()
