# -*- coding: utf-8 -*-
"""
inspect_template.py — Excel 调研模板结构侦察
输出每个工作表的布局/标签/嵌图/图表信息到文本文件, 供建立单元格映射。

用法:
  python inspect_template.py <模板.xlsx> <输出.txt>
"""
import sys
import re
import zipfile
import os

MAX_VAL = 60      # 单元格值截断长度
WIDE_COL_LIMIT = 30  # 列数超过此值视为"数据表", 只详列前几行


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    tpl, out = sys.argv[1], sys.argv[2]

    z = zipfile.ZipFile(tpl)
    names = z.namelist()
    parts = {n: z.read(n) for n in names}
    z.close()

    lines = []

    def w(s=''):
        lines.append(s)

    # ---- 1. 工作表名 -> 文件路径 ----
    wbx = parts['xl/workbook.xml'].decode('utf-8')
    wbrels = parts['xl/_rels/workbook.xml.rels'].decode('utf-8')
    rid_to_target = {}
    for m in re.finditer(r'<Relationship Id="(rId\d+)"[^>]*Target="([^"]+)"', wbrels):
        rid_to_target[m.group(1)] = m.group(2)

    w('========== 工作表清单 ==========')
    sheet_files = {}
    for m in re.finditer(r'<sheet name="([^"]+)"[^>]*r:id="(rId\d+)"', wbx):
        name, rid = m.group(1), m.group(2)
        t = rid_to_target.get(rid, '?')
        if not t.startswith('xl/'):
            t = 'xl/' + t.lstrip('/')
        sheet_files[name] = t
        w(f'{name}  ->  {t}')

    # ---- 2. 逐表网格 ----
    import openpyxl
    wb = openpyxl.load_workbook(tpl)
    for sn in wb.sheetnames:
        ws = wb[sn]
        w()
        w(f'========== [{sn}]  rows={ws.max_row} cols={ws.max_column} ==========')
        w(f'file: {sheet_files.get(sn, "?")}')
        mc = [str(r) for r in ws.merged_cells.ranges]
        if mc:
            w(f'merged: {mc}')
        # 判定是否数据表
        if ws.max_column > WIDE_COL_LIMIT:
            w('(数据表: 仅列出前3行)')
            row_limit = 3
        else:
            row_limit = ws.max_row
        for row in ws.iter_rows(min_row=1, max_row=row_limit):
            vals = []
            for c in row:
                if c.value not in (None, ''):
                    v = str(c.value).replace('\n', '¶')
                    if len(v) > MAX_VAL:
                        v = v[:MAX_VAL] + '…'
                    vals.append(f'{c.coordinate}={v}')
            if vals:
                w('  ' + ' | '.join(vals))
    wb.close()

    # ---- 3. XML层关键信息 ----
    w()
    w('========== XML 层 ==========')
    for n in sorted(parts):
        if not n.endswith('.xml'):
            continue
        x = parts[n].decode('utf-8', errors='ignore')
        if 'worksheet' in n and n.startswith('xl/worksheets/'):
            disp = re.findall(r'<c r="([A-Z]+\d+)"[^>]*>[^<]*<f[^>]*>[^<]*DISPIMG', x)
            hl = re.findall(r'<hyperlink ref="([A-Z]+\d+)" r:id="(rId\d+)"', x)
            dr = re.search(r'<drawing r:id="(rId\d+)"/>', x)
            if disp or hl or dr:
                w(f'[{n}]')
                if disp:
                    w(f'  DISPIMG嵌图: {disp[:40]}{"..." if len(disp) > 40 else ""}')
                if hl:
                    w(f'  hyperlinks: {hl[:20]}{"..." if len(hl) > 20 else ""}')
                if dr:
                    w(f'  drawing引用: {dr.group(1)}')
        elif n.startswith('xl/drawings/') and n.endswith('.xml'):
            anchors = re.findall(r'<xdr:(?:twoCell|oneCell)Anchor', x)
            embeds = re.findall(r'r:embed="(rId\d+)"', x)
            froms = re.findall(r'<xdr:col>(\d+)</xdr:col><xdr:colOff>\d+</xdr:colOff><xdr:row>(\d+)</xdr:row>', x)
            w(f'[{n}] anchors={len(anchors)} embeds={embeds}')
            if froms:
                w(f'  锚点位置(0-based col,row): {froms[:30]}')
        elif n.startswith('xl/charts/'):
            cats = re.findall(r'<c:pt idx="\d+"><c:v>([^<]{0,40})</c:v></c:pt>', x)
            w(f'[{n}] 图表分类缓存: {cats[:12]}')

    w()
    w('========== 其他部件 ==========')
    w('externalLinks: ' + str([n for n in names if 'externalLink' in n and n.endswith('.xml')]))
    w('media数: ' + str(len([n for n in names if n.startswith('xl/media/')])))
    ss = parts['xl/sharedStrings.xml'].decode('utf-8')
    w('sharedStrings条目数: ' + str(len(re.findall(r'<si>', ss))))

    with open(out, 'w', encoding='utf-8-sig') as f:
        f.write('\n'.join(lines))
    print(f'OK 模板结构已输出 -> {out} ({len(lines)}行)')
    print('工作表:', list(sheet_files.keys()))


if __name__ == '__main__':
    main()
